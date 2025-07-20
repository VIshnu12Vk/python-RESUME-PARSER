import streamlit as st
import PyPDF2
import docx
import spacy
from fpdf import FPDF
import io

# Load spaCy model
nlp = spacy.load("en_core_web_sm")

# --- Helper Functions ---
def extract_text_from_pdf(uploaded_file):
    reader = PyPDF2.PdfReader(uploaded_file)
    text = ""
    for page in reader.pages:
        text += page.extract_text()
    return text

def extract_text_from_docx(uploaded_file):
    doc = docx.Document(uploaded_file)
    return "\n".join([para.text for para in doc.paragraphs])

def extract_resume_text(uploaded_file):
    if uploaded_file.name.endswith(".pdf"):
        return extract_text_from_pdf(uploaded_file)
    elif uploaded_file.name.endswith(".docx"):
        return extract_text_from_docx(uploaded_file)
    else:
        return ""

def extract_skills(text):
    doc = nlp(text)
    skills = [ent.text for ent in doc.ents if ent.label_ in ['SKILL', 'ORG', 'WORK_OF_ART']]
    return list(set(skills))



def keyword_match(resume_text, jd_text):
    resume_doc = nlp(resume_text.lower())
    jd_doc = nlp(jd_text.lower())

    resume_tokens = set([token.lemma_ for token in resume_doc if token.is_alpha])
    jd_tokens = set([token.lemma_ for token in jd_doc if token.is_alpha])

    matched = jd_tokens.intersection(resume_tokens)
    score = (len(matched) / len(jd_tokens)) * 100 if jd_tokens else 0

    return round(score, 2), matched

def generate_pdf_summary(name, skills, score, matched_keywords):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    pdf.cell(200, 10, txt="Resume Fit Report", ln=True, align='C')
    pdf.ln(10)
    pdf.cell(200, 10, txt=f"Candidate Name: {name}", ln=True)
    pdf.cell(200, 10, txt=f"Match Score: {score}%", ln=True)
    pdf.ln(5)
    pdf.cell(200, 10, txt="Extracted Skills:", ln=True)
    for skill in skills:
        pdf.cell(200, 10, txt=f"- {skill}", ln=True)

    pdf.ln(5)
    pdf.cell(200, 10, txt="Matched Keywords with Job Description:", ln=True)
    for keyword in matched_keywords:
        pdf.cell(200, 10, txt=f"- {keyword}", ln=True)

    buffer = io.BytesIO()
    pdf.output(buffer)
    buffer.seek(0)
    return buffer

# --- Streamlit GUI ---
st.set_page_config(page_title="Resume Parser + Matcher", layout="centered")
st.title("📄 Resume Parser + Keyword Matcher")

with st.sidebar:
    st.header("Upload Files")
    resume_file = st.file_uploader("Upload Resume (.pdf/.docx)", type=["pdf", "docx"])
    jd_text = st.text_area("Paste Job Description")

if resume_file and jd_text:
    st.subheader("🔍 Analyzing Resume...")

    resume_text = extract_resume_text(resume_file)
    skills = extract_skills(resume_text)
    score, matched_keywords = keyword_match(resume_text, jd_text)

    st.success(f"✅ Match Score: {score}%")
    st.markdown("### 🧠 Extracted Skills")
    st.write(skills)

    st.markdown("### 🔑 Matched Keywords")
    st.write(matched_keywords)

    name = st.text_input("Enter Candidate Name for Report", value="John Doe")

    if st.button("📤 Generate PDF Summary"):
        pdf_buffer = generate_pdf_summary(name, skills, score, matched_keywords)
        st.download_button("⬇️ Download Report", data=pdf_buffer, file_name="resume_summary.pdf", mime="application/pdf")
else:
    st.info("Please upload a resume and paste a job description to proceed.")

